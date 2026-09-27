"""Shared utilities: file extraction, text helpers."""

import os
from io import BytesIO


def extract_text(uploaded_file) -> str:
    """Extract plain text from an uploaded file (PDF, DOCX, TXT, MD).

    Args:
        uploaded_file: A Streamlit UploadedFile object.

    Returns:
        Extracted text as a single string.

    Raises:
        ValueError: If the file type is not supported.
    """
    from docx import Document
    from pypdf import PdfReader

    ext = os.path.splitext(uploaded_file.name)[1].lower()
    file_bytes = uploaded_file.getvalue()

    if ext == ".pdf":
        reader = PdfReader(BytesIO(file_bytes))
        pages = [page.extract_text() or "" for page in reader.pages]
        return "\n\n".join(pages)

    elif ext == ".docx":
        document = Document(BytesIO(file_bytes))
        parts = []
        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                parts.append(paragraph.text.strip())
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(parts)

    elif ext in {".txt", ".md"}:
        return file_bytes.decode("utf-8", errors="replace")

    else:
        raise ValueError(f"Unsupported file type: {ext}")
