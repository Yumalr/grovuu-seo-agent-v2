"""XLSX meta data sheet renderer."""

import io
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill


META_COLUMNS = [
    "#",
    "Page Name",
    "URL",
    "Meta Title",
    "Meta Description",
    "H1 (On-Page)",
    "Primary Keywords",
    "Search Intent",
]


def render_excel_meta(meta_rows: list) -> bytes:
    """Render the meta data sheet as an XLSX file.

    Args:
        meta_rows: List of dicts, each with keys matching the column schema.

    Returns:
        XLSX file content as bytes.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Meta Data - Final"
    ws.append(META_COLUMNS)

    # Format headers
    header_fill = PatternFill(
        start_color="1F4E78", end_color="1F4E78", fill_type="solid"
    )
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for i, row in enumerate(meta_rows, start=1):
        pk = row.get("primary_keywords", "")
        if isinstance(pk, list):
            pk = ", ".join(pk)
        ws.append(
            [
                i,
                row.get("page_name", ""),
                row.get("url", ""),
                row.get("meta_title", ""),
                row.get("meta_description", ""),
                row.get("h1", ""),
                pk,
                row.get("search_intent", ""),
            ]
        )

    # Auto-adjust column widths and wrap text
    for col in ws.columns:
        max_length = 0
        column_letter = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except Exception:
                pass
            cell.alignment = Alignment(wrap_text=True, vertical="top")

        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
