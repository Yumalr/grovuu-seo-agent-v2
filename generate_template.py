import os
from docx import Document

doc = Document()
doc.add_heading('SEO Strategy & Website Architecture', 0)

doc.add_heading('What We Are Doing', level=1)
doc.add_paragraph('{{ what_we_are_doing }}')

doc.add_heading('Why This is Necessary', level=1)
doc.add_paragraph('{{ why_necessary }}')

doc.add_heading('How This Plan Aligns With Your Business Goals', level=1)
doc.add_paragraph('{% for goal in business_goals %}')
doc.add_paragraph('Business Driver: {{ goal.driver }}')
doc.add_paragraph('How SEO Achieves This: {{ goal.how_seo_achieves_this }}')
doc.add_paragraph('')
doc.add_paragraph('{% endfor %}')

doc.add_heading('The Value of This Structure: Why a "Flat" Website Fails', level=1)
doc.add_paragraph('{{ structure_value }}')

doc.add_heading('Ideal Customer Profiles', level=1)
doc.add_paragraph('{% for icp in icps %}')
doc.add_paragraph('- {{ icp }}')
doc.add_paragraph('{% endfor %}')

doc.add_heading('IA & URL Plan', level=1)
doc.add_paragraph('{% for page in pages %}')
doc.add_paragraph('Type: {{ page.type }}')
doc.add_paragraph('Proposed URL: {{ page.proposed_url }}')
doc.add_paragraph('Main Keyword: {{ page.main_keyword }}')
doc.add_paragraph('Title (Patterned): {{ page.title_patterned }}')
doc.add_paragraph('H1 (Patterned): {{ page.h1_patterned }}')
doc.add_paragraph('Phase: {{ page.phase }} | Policy: {{ page.policy }}')
doc.add_paragraph('')
doc.add_paragraph('{% endfor %}')

doc.add_heading('Next Steps', level=1)
doc.add_paragraph('{% for step in next_steps %}')
doc.add_paragraph('- {{ step }}')
doc.add_paragraph('{% endfor %}')

doc.save('template.docx')
print("Template created successfully.")
